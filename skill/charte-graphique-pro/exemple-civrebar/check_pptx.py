"""Contrôle de débordement du texte dans un .pptx, sans LibreOffice :
mesure chaque ligne avec les vraies polices (fontTools) et compare à la zone de texte."""
import sys
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Emu

FONTS = "/home/user/Claude/logo/fonts/"
FILES = {("Sora", True): "sora-latin-700-normal.woff2", ("Sora", False): "sora-latin-400-normal.woff2",
         ("Manrope", False): "manrope-latin-400-normal.woff2", ("Manrope", True): "manrope-latin-700-normal.woff2",
         ("Instrument Serif", False): "instrument-serif-latin-400-italic.woff2",
         ("JetBrains Mono", False): "jetbrains-mono-latin-400-normal.woff2"}
_cache = {}


def advance(face, bold, text, size_pt, spacing_pt=0.0):
    key = (face, bold) if (face, bold) in FILES else (face, False)
    if key not in _cache:
        f = TTFont(FONTS + FILES[key])
        _cache[key] = (f.getBestCmap(), f["hmtx"], f["head"].unitsPerEm)
    cmap, hmtx, upm = _cache[key]
    w = sum(hmtx[cmap.get(ord(c), cmap[ord("x")])][0] for c in text) * size_pt / upm
    return w + spacing_pt * len(text)


def lines_needed(face, bold, text, size, width_pt, spacing):
    n = 0
    for para in text.split("\n"):
        cur, n_para = "", 1
        for word in para.split(" "):
            trial = (cur + " " + word).strip()
            if advance(face, bold, trial, size, spacing) > width_pt and cur:
                n_para += 1
                cur = word
            else:
                cur = trial
        n += n_para
    return n


def main(path):
    bad = 0
    for i, slide in enumerate(Presentation(path).slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            run = next(r for p in sh.text_frame.paragraphs for r in p.runs)
            face, size = run.font.name, run.font.size.pt
            bold = bool(run.font.bold)
            rpr = run._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
            spc = int(rpr.get("spc", 0)) / 100 if rpr is not None else 0
            tf = sh.text_frame
            ml = (tf.margin_left or 0) + (tf.margin_right or 0)
            mt = (tf.margin_top or 0) + (tf.margin_bottom or 0)
            w_pt = Emu(sh.width - ml).pt
            h_pt = Emu(sh.height - mt).pt
            text = sh.text_frame.text
            n = lines_needed(face, bold, text, size, w_pt, spc)
            need = n * size * 1.2
            longest = max(advance(face, bold, w, size, spc) for w in text.replace("\n", " ").split(" "))
            flag = need > h_pt + 1 or longest > w_pt
            bad += flag
            if flag or "-v" in sys.argv:
                print(f"{'!!' if flag else 'ok'} diapo {i}: {text[:40]!r} lignes={n} besoin={need:.0f}pt dispo={h_pt:.0f}pt")
    print("débordements :", bad)


if __name__ == "__main__":
    main(sys.argv[1])
