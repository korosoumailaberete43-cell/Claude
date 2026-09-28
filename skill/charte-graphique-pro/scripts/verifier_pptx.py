"""Détecte le texte qui déborde de sa zone dans un .pptx, sans LibreOffice.
Mesure chaque ligne avec les vrais fichiers de police (fontTools) et compare à la boîte.

    python3 verifier_pptx.py deck.pptx dossier_polices/ [-v]

Le dossier contient les polices utilisées (.ttf/.otf/.woff2). Le fichier est choisi par nom :
la famille (« Sora », « JetBrains Mono »…) doit apparaître dans le nom du fichier, et les
fichiers « 700 »/« bold » servent au gras, « italic » à l'italique. Faute de fichier trouvé,
la police est signalée et la zone ignorée.
Dépendances : pip install python-pptx fonttools brotli
"""
import pathlib
import re
import sys

from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_polices, _cache = [], {}


def trouver(famille, gras, italique):
    cle = re.sub(r"[^a-z]", "", famille.lower())
    cands = [p for p in _polices if cle in re.sub(r"[^a-z]", "", p.stem.lower())]
    if not cands:
        return None

    def score(p):
        n = p.stem.lower()
        s = 0
        s += 2 if (("700" in n or "bold" in n) == gras) else 0
        s += 1 if (("italic" in n) == italique) else 0
        s += 0.5 if ("400" in n or "regular" in n) and not gras else 0
        return s
    return max(cands, key=score)


def avance(fichier, texte, taille, espacement):
    if fichier not in _cache:
        f = TTFont(fichier)
        _cache[fichier] = (f.getBestCmap(), f["hmtx"], f["head"].unitsPerEm)
    cmap, hmtx, upm = _cache[fichier]
    repli = cmap.get(ord("x"))
    return sum(hmtx[cmap.get(ord(c), repli)][0] for c in texte) * taille / upm + espacement * len(texte)


def lignes(fichier, texte, taille, largeur, esp):
    n = 0
    for para in texte.replace("\x0b", "\n").split("\n"):
        cur, k = "", 1
        for mot in para.split(" "):
            essai = (cur + " " + mot).strip()
            if avance(fichier, essai, taille, esp) > largeur and cur:
                k, cur = k + 1, mot
            else:
                cur = essai
        n += k
    return n


def main(chemin, dossier, bavard=False):
    _polices.extend(p for p in pathlib.Path(dossier).iterdir() if p.suffix.lower() in (".ttf", ".otf", ".woff2"))
    defauts = 0
    for i, diapo in enumerate(Presentation(chemin).slides, 1):
        for sh in diapo.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            run = next((r for p in sh.text_frame.paragraphs for r in p.runs), None)
            if run is None or run.font.size is None:
                continue
            f = trouver(run.font.name or "", bool(run.font.bold), bool(run.font.italic))
            if f is None:
                print(f"?? diapo {i}: police « {run.font.name} » introuvable dans {dossier}")
                continue
            rpr = run._r.find(f"{A}rPr")
            esp = int(rpr.get("spc", 0)) / 100 if rpr is not None else 0
            tf = sh.text_frame
            larg = Emu(sh.width - (tf.margin_left or 0) - (tf.margin_right or 0)).pt
            haut = Emu(sh.height - (tf.margin_top or 0) - (tf.margin_bottom or 0)).pt
            taille = run.font.size.pt
            n = lignes(f, tf.text, taille, larg, esp)
            besoin = n * taille * 1.2
            mot_max = max(avance(f, m, taille, esp) for m in tf.text.split())
            ko = besoin > haut + 1 or mot_max > larg
            defauts += ko
            if ko or bavard:
                print(f"{'!!' if ko else 'ok'} diapo {i}: {tf.text[:40]!r} lignes={n} besoin={besoin:.0f}pt dispo={haut:.0f}pt")
    print("débordements :", defauts)
    return defauts


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1], sys.argv[2], "-v" in sys.argv) else 0)
