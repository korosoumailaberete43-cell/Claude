"""Aperçu approximatif d'un .pptx simple (textes, formes, images, fonds) via HTML + Chromium.
Remplace LibreOffice, indisponible ici. python3 preview_pptx.py deck.pptx sortie_dir"""
import base64
import os
import pathlib
import subprocess
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from brand import font_css  # polices Step by Step (Quicksand, Comfortaa, Nunito)

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PX = 96  # px par pouce


def px(v):
    return Emu(v).inches * PX


P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def color_of(el, tag):
    ns = P if tag == "spPr" else A
    n = el.find(f".//{ns}{tag}/{A}solidFill/{A}srgbClr")
    return "#" + n.get("val") if n is not None else None


FAM = {}


def slide_html(prs, slide, pptx_path):
    W, H = px(prs.slide_width), px(prs.slide_height)
    bg = ""
    cs = slide._element.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}bg")
    if cs is not None:
        blip = cs.find(f".//{A}blip")
        clr = cs.find(f".//{A}srgbClr")
        if blip is not None:
            part = slide.part.related_part(blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"))
            bg = f"background:url(data:image/png;base64,{base64.b64encode(part.blob).decode()}) 0 0/100% 100%;"
        elif clr is not None:
            bg = f"background:#{clr.get('val')};"
    out = []
    for sh in slide.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        box = f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            out.append(f'<img style="{box}" src="data:image/png;base64,{base64.b64encode(sh.image.blob).decode()}">')
            continue
        sp = sh._element
        prst = sp.find(f".//{A}prstGeom")
        geom = prst.get("prst") if prst is not None else "rect"
        fill = color_of(sp, "spPr")
        ln = sp.find(f".//{A}ln")
        line = color_of(sp, "ln") if ln is not None else None
        lw = int(ln.get("w", 12700)) / 12700 * 96 / 72 if ln is not None else 0
        if geom == "line":
            out.append(f'<div style="position:absolute;left:{x}px;top:{y - lw / 2}px;width:{w}px;height:{lw}px;background:{line}"></div>')
        elif not sh.has_text_frame or not sh.text_frame.text.strip() or fill:
            rad = "50%" if geom == "ellipse" else ("8px" if geom == "roundRect" else "0")
            b = f"border:{lw}px solid {line};" if line else ""
            out.append(f'<div style="{box}background:{fill or "transparent"};border-radius:{rad};{b}box-sizing:border-box"></div>')
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            run = next(r for p in tf.paragraphs for r in p.runs)
            f = run.font
            col = "#" + str(f.color.rgb) if f.color and f.color.type else "#000"
            rpr = run._r.find(f"{A}rPr")
            spc = int(rpr.get("spc", 0)) / 100 if rpr is not None else 0
            al = tf.paragraphs[0].alignment
            align = {1: "left", 2: "center", 3: "right"}.get(int(al) if al else 1, "left")
            anchor = tf._txBody.find(f"{A}bodyPr").get("anchor", "t")
            jc = {"t": "flex-start", "ctr": "center", "b": "flex-end"}[anchor]
            m = [px(v or 0) for v in (tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left)]
            txt = tf.text.replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br>").replace("\x0b", "<br>")
            out.append(
                f'<div style="{box}display:flex;flex-direction:column;justify-content:{jc};padding:{m[0]}px {m[1]}px {m[2]}px {m[3]}px;'
                f'box-sizing:border-box;outline:1px dashed rgba(255,0,150,.0)">'
                f'<div style="font-family:{FAM.get(f.name, f.name)};font-size:{f.size.pt * 96 / 72}px;line-height:1.2;'
                f'font-weight:{700 if f.bold else 400};font-style:{"italic" if f.italic else "normal"};color:{col};'
                f'letter-spacing:{spc * 96 / 72}px;text-align:{align}">{txt}</div></div>')
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}*{{margin:0}}'
            f'body{{width:{W}px;height:{H}px;position:relative;overflow:hidden;{bg}}}</style></head><body>{"".join(out)}</body></html>')


def main(pptx_path, out_dir):
    out = pathlib.Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(pptx_path)
    W, H = round(px(prs.slide_width)), round(px(prs.slide_height))
    jobs = []
    for i, s in enumerate(prs.slides, 1):
        f = out / f"slide-{i:02d}.html"
        f.write_text(slide_html(prs, s, pptx_path), encoding="utf-8")
        jobs.append(f"{f}|{out / f'slide-{i:02d}.png'}|{W}|{H}")
    env = dict(os.environ, NODE_PATH=subprocess.check_output(["npm", "root", "-g"], text=True).strip())
    subprocess.run(["node", str(pathlib.Path(__file__).parent / "render_png.js"), *jobs], check=True, env=env)
    print("ok", len(jobs))


if __name__ == "__main__":
    main(*sys.argv[1:3])
